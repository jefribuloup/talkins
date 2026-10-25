"""Pembuat slide dari data Excel (gaya Gamma / PowerPoint).

Alur:
1. build_deck()  : susun slide (judul, kartu angka, grafik, tabel, kualitas data, penutup) dari hasil
                   recap_engine — angka & grafik SELALU dari data asli, bukan karangan AI.
2. polish_deck() : (opsional, panggilan terpisah) minta AI menulis ulang judul & poin tiap slide
                   supaya berbentuk insight; gagal -> teks bawaan dipakai.
3. render_pptx() : ekspor ke file .pptx sungguhan (grafik native, bisa diedit di PowerPoint).
"""

import io
import json
import re
from datetime import date

from lib import db
from lib import excel_engine

THEMES = {
    "aurora": {"name": "Aurora", "bg": "0B1220", "bg2": "162B4D", "fg": "F1F5F9", "muted": "94A3B8",
               "accent": "38BDF8", "accent2": "A78BFA", "card": "16233B",
               "palette": ["38BDF8", "A78BFA", "34D399", "FBBF24", "F472B6", "60A5FA", "FB923C", "2DD4BF"]},
    "sunrise": {"name": "Sunrise", "bg": "FFF7ED", "bg2": "FFEDD5", "fg": "1F2937", "muted": "6B7280",
                "accent": "F97316", "accent2": "E11D48", "card": "FFFFFF",
                "palette": ["F97316", "E11D48", "0EA5E9", "16A34A", "8B5CF6", "EAB308", "14B8A6", "EC4899"]},
    "forest": {"name": "Forest", "bg": "0F2A1D", "bg2": "1B4332", "fg": "ECFDF5", "muted": "A7C4B5",
               "accent": "84CC16", "accent2": "22D3EE", "card": "17382A",
               "palette": ["84CC16", "22D3EE", "FACC15", "FB7185", "A78BFA", "34D399", "F97316", "60A5FA"]},
    "mono": {"name": "Mono", "bg": "FFFFFF", "bg2": "F1F5F9", "fg": "0F172A", "muted": "64748B",
             "accent": "2563EB", "accent2": "0F172A", "card": "F8FAFC",
             "palette": ["2563EB", "0F172A", "64748B", "F59E0B", "10B981", "EF4444", "8B5CF6", "06B6D4"]},
}

T = {
    "id": {
        "title_sub": "{type} · {rows} baris × {cols} kolom · sheet {sheet}",
        "kpi.title": "Sekilas data", "kpi.headline": "Gambaran umum data",
        "cat.title": "Komposisi {col}", "cat.sum": "{m} per {col}", "cat.cap1": "{v} terbesar: {n} data ({pct}%)",
        "cat.cap2": "{k} kelompok teratas mencakup {pct}% dari seluruh data",
        "cat.capsum": "{v} memberi kontribusi terbesar: {n} ({pct}%)",
        "hist.title": "Sebaran {col}", "hist.cap": "Nilai {col} berkisar {lo} sampai {hi}; paling banyak di rentang {a}–{b}",
        "trend.title": "Tren {col} per {gran}", "trend.month": "bulan", "trend.day": "hari",
        "trend.cap": "Puncak di {peak} ({val}); titik terendah di {low} ({lv})",
        "table.title": "Ringkasan per {col}", "table.cap": "Tabel merangkum jumlah data dan nilai per kelompok",
        "col.count": "Jumlah", "col.sum": "Total {m}", "col.avg": "Rata-rata {m}", "col.others": "Lainnya", "col.total": "Total",
        "q.title": "Kualitas data", "q.cap": "Skor kesehatan data {score}/100 ({grade})",
        "q.ok": "Tidak ada masalah berarti ditemukan.",
        "end.title": "Kesimpulan & langkah berikutnya",
        "end.fix": "Perbaiki: {t}", "end.use": "Gunakan temuan di atas sebagai dasar laporan atau keputusan.",
        "end.pivot": "Rangkum per {col} untuk melihat kelompok yang paling berpengaruh.",
        "sev.high": "Tinggi", "sev.medium": "Sedang", "sev.low": "Rendah", "sev.info": "Info", "sev.ok": "Aman",
        "deck.default": "Ringkasan {name}",
    },
    "en": {
        "title_sub": "{type} · {rows} rows × {cols} columns · sheet {sheet}",
        "kpi.title": "Data at a glance", "kpi.headline": "Overview of the data",
        "cat.title": "Breakdown of {col}", "cat.sum": "{m} by {col}", "cat.cap1": "{v} is largest: {n} records ({pct}%)",
        "cat.cap2": "The top {k} groups cover {pct}% of all records",
        "cat.capsum": "{v} contributes most: {n} ({pct}%)",
        "hist.title": "Distribution of {col}", "hist.cap": "{col} ranges from {lo} to {hi}; most values fall in {a}–{b}",
        "trend.title": "{col} trend by {gran}", "trend.month": "month", "trend.day": "day",
        "trend.cap": "Peak in {peak} ({val}); lowest in {low} ({lv})",
        "table.title": "Summary by {col}", "table.cap": "The table summarises record counts and values per group",
        "col.count": "Count", "col.sum": "Total {m}", "col.avg": "Avg {m}", "col.others": "Others", "col.total": "Total",
        "q.title": "Data quality", "q.cap": "Data health score {score}/100 ({grade})",
        "q.ok": "No significant issues found.",
        "end.title": "Conclusions & next steps",
        "end.fix": "Fix: {t}", "end.use": "Use the findings above as the basis for reports or decisions.",
        "end.pivot": "Summarise by {col} to see which groups matter most.",
        "sev.high": "High", "sev.medium": "Medium", "sev.low": "Low", "sev.info": "Info", "sev.ok": "OK",
        "deck.default": "{name} summary",
    },
}


def _fmt(x, lang="id"):
    if x is None:
        return "–"
    if isinstance(x, str):
        return x
    if abs(x) >= 1000:
        s = f"{x:,.0f}"
    elif x == int(x):
        s = str(int(x))
    else:
        s = f"{x:,.2f}".rstrip("0").rstrip(".")
    return s.replace(",", "§").replace(".", ",").replace("§", ".") if lang == "id" else s


def _short(s, n):
    s = str(s)
    return s if len(s) <= n else s[: n - 1] + "…"


def build_deck(file_id, spec):
    spec = spec or {}
    lang = spec.get("lang") if spec.get("lang") in T else "id"
    X = T[lang]
    sheets = db.list_sheet_names(file_id)
    if not sheets:
        raise ValueError("File ini belum punya sheet yang bisa dibuat slide.")
    sheet = spec.get("sheet") if spec.get("sheet") in sheets else sheets[0]
    theme = spec.get("theme") if spec.get("theme") in THEMES else "aurora"
    fmeta = db.get_file(file_id) or {}
    name = re.sub(r"\.(xlsx|xls|xlsm)$", "", fmeta.get("filename") or "data", flags=re.I)
    name = re.sub(r"[_\s]+", " ", name).strip() or "data"

    r = excel_engine.get_recap(file_id, sheet, lang)
    ds, summ = r["dataset"], r["summary"]
    topic = (spec.get("topic") or "").strip()
    slides = []
    names = {x["col"]: x["label"] for x in ds["roles"] if x.get("label")}

    def dn(col):  # nama kolom untuk judul: label kolom kalau ada, kalau tidak huruf kolomnya
        return _short(names.get(col) or col, 24)

    def add(sl):
        sl["id"] = f"s{len(slides) + 1}"
        slides.append(sl)

    add({"type": "title", "title": topic or X["deck.default"].format(name=name),
         "subtitle": X["title_sub"].format(type=ds["label"], rows=_fmt(summ["rows"], lang),
                                            cols=summ["cols"], sheet=sheet),
         "chips": [ds["label"], f"{summ['fill_pct']}%", date.today().strftime("%d/%m/%Y")]})

    if ds["kpis"]:
        add({"type": "kpis", "title": X["kpi.headline"],
             "items": [{"label": k["label"], "value": _fmt(k["value"], lang), "sub": k.get("sub", "")}
                       for k in ds["kpis"][:6]]})

    # Grafik komposisi kategori (pakai jumlah nilai uang/kuantitas kalau ada, selain itu jumlah data)
    for ci, cat in enumerate(r["categories"][:2]):
        pv = next((p for p in r["pivots"] if p["col"] == cat["col"]), None)
        meas = pv["measures"][0] if pv and pv["measures"] else None
        labels, values, caption = None, None, ""
        if pv and meas and meas["role"] in ("money", "quantity", "measure"):
            rows = [x for x in pv["rows"] if x["label"] is not None]
            tot = sum(x["m"][0]["sum"] for x in rows) or 1
            labels = [x["label"] for x in rows]
            values = [x["m"][0]["sum"] for x in rows]
            best = max(rows, key=lambda x: x["m"][0]["sum"])
            title = X["cat.sum"].format(m=dn(meas["col"]), col=dn(cat["col"]))
            caption = X["cat.capsum"].format(v=best["label"], n=_fmt(best["m"][0]["sum"], lang),
                                             pct=round(100 * best["m"][0]["sum"] / tot))
        else:
            labels = [x["value"] for x in cat["top"]]
            values = [x["count"] for x in cat["top"]]
            if cat["others"] > 0:
                labels.append(X["col.others"])
                values.append(cat["others"])
            tot = sum(values) or 1
            title = X["cat.title"].format(col=dn(cat["col"]))
            caption = X["cat.cap1"].format(v=cat["top"][0]["value"], n=cat["top"][0]["count"],
                                           pct=round(100 * cat["top"][0]["count"] / tot))
        bullets = [caption]
        if len(values) >= 3:
            k = min(3, len(values))
            share = round(100 * sum(sorted(values, reverse=True)[:k]) / (sum(values) or 1))
            bullets.append(X["cat.cap2"].format(k=k, pct=share))
        add({"type": "chart", "kind": "pie" if len(labels) <= 6 else "column", "title": title,
             "bullets": bullets, "labels": [_short(x, 22) for x in labels],
             "series": [{"name": "", "values": values}]})

    if r["pivots"]:
        pv = r["pivots"][0]
        cols = [dn(pv["col"]), X["col.count"]]
        kinds = ["text", "num"]
        bars = [1]  # kolom yang diberi data bar di slide: jumlah data & total (rata-rata tidak)
        for m in pv["measures"]:
            if m["role"] not in ("score", "percent"):
                cols.append(X["col.sum"].format(m=dn(m["col"])))
                kinds.append("num")
                bars.append(len(cols) - 1)
            cols.append(X["col.avg"].format(m=dn(m["col"])))
            kinds.append("num")

        def row_cells(rw, label):
            cells, raw = [label, _fmt(rw["count"], lang)], [None, rw["count"]]
            for j, m in enumerate(pv["measures"]):
                if m["role"] not in ("score", "percent"):
                    cells.append(_fmt(rw["m"][j]["sum"], lang))
                    raw.append(rw["m"][j]["sum"])
                cells.append(_fmt(rw["m"][j]["avg"], lang))
                raw.append(rw["m"][j]["avg"])
            return cells, raw

        trs, raws = [], []
        for rw in pv["rows"][:8]:
            c_, w_ = row_cells(rw, _short(rw["label"] if rw["label"] is not None else X["col.others"], 22))
            trs.append(c_)
            raws.append(w_)
        c_, w_ = row_cells(pv["total"], X["col.total"])
        add({"type": "table", "title": X["table.title"].format(col=dn(pv["col"])), "bullets": [X["table.cap"]],
             "columns": cols, "kinds": kinds, "bars": bars, "rows": trs, "raw": raws, "total": c_})

    if r["trend"]:
        tr = r["trend"]
        pts = tr["points"]
        s0 = [p["values"][0] for p in pts]
        hi_i, lo_i = s0.index(max(s0)), s0.index(min(s0))
        add({"type": "chart", "kind": "line",
             "title": X["trend.title"].format(col=" / ".join(dn(c) for c in tr["cols"]), gran=X["trend." + tr["granularity"]]),
             "bullets": [X["trend.cap"].format(peak=pts[hi_i]["label"], val=_fmt(s0[hi_i], lang),
                                               low=pts[lo_i]["label"], lv=_fmt(s0[lo_i], lang))],
             "labels": [p["label"] for p in pts],
             "series": [{"name": dn(nm), "values": [p["values"][i] for p in pts]} for i, nm in enumerate(tr["cols"])]})

    if r["histograms"]:
        hg = r["histograms"][0]
        bins = hg["bins"]
        mx = max(range(len(bins)), key=lambda i: bins[i]["count"])
        add({"type": "chart", "kind": "column",
             "title": X["hist.title"].format(col=dn(hg["col"])),
             "bullets": [X["hist.cap"].format(col=dn(hg["col"]), lo=_fmt(bins[0]["lo"], lang), hi=_fmt(bins[-1]["hi"], lang),
                                              a=_fmt(bins[mx]["lo"], lang), b=_fmt(bins[mx]["hi"], lang))],
             "mono": True, "label_end": _fmt(bins[-1]["hi"], lang), "labels": [_fmt(b["lo"], lang) for b in bins],
             "series": [{"name": "", "values": [b["count"] for b in bins]}]})

    issues = [f for f in r["findings"] if f["category"] not in ("insight",)][:5]
    add({"type": "findings", "title": X["q.title"], "score": summ["score"], "grade": summ["grade"],
         "bullets": [X["q.cap"].format(score=summ["score"], grade=summ["grade"])],
         "items": [{"severity": f["severity"], "label": X["sev." + f["severity"]], "title": f["title"],
                    "detail": _short(f["detail"], 150)} for f in issues]
         or [{"severity": "ok", "label": X["sev.ok"], "title": X["q.ok"], "detail": ""}]})

    end = []
    for f in [f for f in r["findings"] if f["category"] not in ("insight", "ok")][:2]:
        end.append(X["end.fix"].format(t=f["title"]))
    if r["pivots"]:
        end.append(X["end.pivot"].format(col=dn(r["pivots"][0]["col"])))
    end.append(X["end.use"])
    add({"type": "closing", "title": X["end.title"], "bullets": end[:4]})

    return {"title": slides[0]["title"], "subtitle": slides[0]["subtitle"], "lang": lang, "sheet": sheet,
            "theme": theme, "themes": THEMES, "slides": slides}


# ---------------------------------------------------------------- AI polish

def polish_deck(deck):
    """Minta AI menulis ulang judul & poin tiap slide jadi kalimat insight (gaya Gamma).
    Hanya teks; angka/grafik tidak disentuh. Return {"slides": {id: {title, bullets}}} atau {}."""
    from lib import chat_engine
    lang = deck.get("lang", "id")
    facts = []
    for s in deck["slides"]:
        f = {"id": s["id"], "type": s["type"], "title": s.get("title", "")}
        if s["type"] == "chart":
            f["labels"] = s["labels"][:8]
            f["values"] = [round(v, 2) for v in s["series"][0]["values"][:8]]
        elif s["type"] == "kpis":
            f["items"] = [f"{i['label']}: {i['value']}" for i in s["items"]]
        elif s["type"] == "table":
            f["columns"] = s["columns"]
            f["rows"] = s["rows"][:5]
        elif s["type"] == "findings":
            f["items"] = [f"{i['label']}: {i['title']}" for i in s["items"]]
        elif s["type"] == "closing":
            f["draft"] = s["bullets"]
        facts.append(f)
    language = "Bahasa Indonesia" if lang == "id" else "English"
    prompt = (
        f"Tulis ulang teks slide presentasi berikut dalam {language}. Gaya: insight yang tegas dan ringkas "
        "seperti slide konsultan. Untuk SETIAP slide (kecuali type 'title') beri 'title' (maks 12 kata, berupa "
        "kesimpulan, bukan label) dan 'bullets' (2-3 poin, maks 18 kata/poin, hanya memakai fakta yang diberikan, "
        "jangan mengarang angka). Slide 'closing' berisi rekomendasi tindakan konkret. Balas HANYA JSON: "
        '{"slides": {"<id>": {"title": "...", "bullets": ["..."]}}}\n\n'
        f"Topik/arahan user: {deck.get('title')}\nData slide:\n{json.dumps(facts, ensure_ascii=False)}"
    )
    text, _ = chat_engine._call_llm(
        "Kamu penulis slide data yang akurat. Selalu balas JSON valid saja, tanpa markdown.", [], prompt)
    cleaned = text.strip().strip("`")
    if cleaned.lower().startswith("json"):
        cleaned = cleaned[4:].strip()
    for cand in (cleaned, cleaned[cleaned.find("{"): cleaned.rfind("}") + 1]):
        try:
            data = json.loads(cand)
            break
        except (json.JSONDecodeError, ValueError):
            data = None
    out = {}
    if isinstance(data, dict):
        for sid, v in (data.get("slides") or {}).items():
            if isinstance(v, dict):
                title = str(v.get("title") or "").strip()[:140]
                bullets = [str(b).strip()[:200] for b in (v.get("bullets") or []) if str(b).strip()][:4]
                if title or bullets:
                    out[sid] = {"title": title, "bullets": bullets}
    return {"slides": out}


# ---------------------------------------------------------------- PPTX

def _rgb(h):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(h)


def render_pptx(deck, theme_key=None):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt

    th = THEMES.get(theme_key or deck.get("theme")) or THEMES["aurora"]
    pal = th["palette"]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    sev_color = {"high": "EF4444", "medium": "F59E0B", "low": "38BDF8", "info": "38BDF8", "ok": "22C55E"}

    def bg(slide, color):
        f = slide.background.fill
        f.solid()
        f.fore_color.rgb = _rgb(color)

    def rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE, line=None, round_=None):
        sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.fill.solid()
        sh.fill.fore_color.rgb = _rgb(color)
        if line:
            sh.line.color.rgb = _rgb(line)
            sh.line.width = Pt(0.75)
        else:
            sh.line.fill.background()
        if round_ is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            sh.adjustments[0] = round_
        sh.shadow.inherit = False
        return sh

    def text(slide, s, x, y, w, h, size=18, bold=False, color=None, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, font="Calibri"):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.05)
        p = tf.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = str(s)
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = font
        run.font.color.rgb = _rgb(color or th["fg"])
        return tb

    def heading(slide, title):
        rect(slide, 0.6, 0.62, 0.12, 0.62, th["accent"])
        text(slide, title, 0.85, 0.5, 11.8, 0.9, size=30, bold=True, anchor=MSO_ANCHOR.MIDDLE)

    def bullets_box(slide, items, x, y, w, h, size=18):
        rect(slide, x, y, w, h, th["card"], MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.05)
        tb = slide.shapes.add_textbox(Inches(x + 0.3), Inches(y + 0.3), Inches(w - 0.6), Inches(h - 0.6))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, b in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(14)
            r1 = p.add_run()
            r1.text = "● "
            r1.font.size = Pt(size - 4)
            r1.font.color.rgb = _rgb(th["accent"])
            r2 = p.add_run()
            r2.text = str(b)
            r2.font.size = Pt(size)
            r2.font.name = "Calibri"
            r2.font.color.rgb = _rgb(th["fg"])

    for sl in deck["slides"]:
        s = prs.slides.add_slide(blank)
        bg(s, th["bg"])
        t = sl["type"]
        if t == "title":
            rect(s, 8.6, -1.4, 7.0, 7.0, th["bg2"], MSO_SHAPE.OVAL)
            rect(s, 10.2, 4.3, 4.4, 4.4, th["accent2"], MSO_SHAPE.OVAL)
            text(s, sl["title"], 0.85, 1.7, 9.2, 2.6, size=46, bold=True, anchor=MSO_ANCHOR.BOTTOM)
            rect(s, 0.9, 4.5, 1.4, 0.14, th["accent"])
            text(s, sl["subtitle"], 0.85, 4.8, 9.0, 0.9, size=18, color=th["muted"])
            x = 0.9
            for chip in sl.get("chips", [])[:3]:
                w = 0.28 + 0.115 * len(chip)
                rect(s, x, 5.75, w, 0.45, th["card"], MSO_SHAPE.ROUNDED_RECTANGLE, line=th["accent"], round_=0.5)
                text(s, chip, x, 5.75, w, 0.45, size=12, color=th["accent"], align=PP_ALIGN.CENTER,
                     anchor=MSO_ANCHOR.MIDDLE)
                x += w + 0.2
        elif t == "kpis":
            heading(s, sl["title"])
            items = sl["items"][:6]
            cols = 3 if len(items) > 4 else 2 if len(items) == 4 else len(items) or 1
            rows_n = -(-len(items) // cols)
            cw, ch = (12.1 - 0.4 * (cols - 1)) / cols, (5.2 - 0.4 * (rows_n - 1)) / rows_n
            for i, it in enumerate(items):
                x = 0.6 + (i % cols) * (cw + 0.4)
                y = 1.7 + (i // cols) * (ch + 0.4)
                rect(s, x, y, cw, ch, th["card"], MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.08)
                rect(s, x, y + 0.3, 0.09, ch - 0.6, pal[i % len(pal)])
                val = str(it["value"])
                text(s, val, x + 0.35, y + 0.25, cw - 0.5, ch * 0.5, size=36 if len(val) < 12 else 24, bold=True,
                     color=pal[i % len(pal)], anchor=MSO_ANCHOR.MIDDLE)
                text(s, it["label"], x + 0.35, y + ch * 0.55, cw - 0.5, 0.5, size=15, color=th["fg"])
                if it.get("sub"):
                    text(s, it["sub"], x + 0.35, y + ch * 0.55 + 0.45, cw - 0.5, 0.45, size=12, color=th["muted"])
        elif t == "chart":
            heading(s, sl["title"])
            kind = sl["kind"]
            cd = CategoryChartData()
            cd.categories = sl["labels"]
            for se in sl["series"]:
                cd.add_series(se["name"] or sl["title"], se["values"])
            ctype = {"pie": XL_CHART_TYPE.DOUGHNUT, "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
                     "bar": XL_CHART_TYPE.BAR_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS}[kind]
            rect(s, 0.6, 1.7, 7.9, 5.2, th["card"], MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.04)
            gf = s.shapes.add_chart(ctype, Inches(0.8), Inches(1.85), Inches(7.5), Inches(4.9), cd)
            ch = gf.chart
            ch.font.size = Pt(12)
            ch.font.color.rgb = _rgb(th["fg"])
            ch.has_title = False
            plot = ch.plots[0]
            if kind == "pie":
                ch.has_legend = True
                ch.legend.position = XL_LEGEND_POSITION.RIGHT
                ch.legend.include_in_layout = False
                ch.legend.font.color.rgb = _rgb(th["fg"])
                ser = plot.series[0]
                for i in range(len(sl["labels"])):
                    pt = ser.points[i]
                    pt.format.fill.solid()
                    pt.format.fill.fore_color.rgb = _rgb(pal[i % len(pal)])
                plot.has_data_labels = True
                plot.data_labels.font.size = Pt(11)
                plot.data_labels.font.color.rgb = _rgb(th["fg"] if th["bg"] != "FFFFFF" else "FFFFFF")
                plot.data_labels.show_percentage = True
                plot.data_labels.show_value = False
            else:
                ch.has_legend = len(sl["series"]) > 1
                if ch.has_legend:
                    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
                    ch.legend.include_in_layout = False
                    ch.legend.font.color.rgb = _rgb(th["fg"])
                va = ch.value_axis
                va.major_gridlines.format.line.color.rgb = _rgb(th["muted"])
                va.major_gridlines.format.line.width = Pt(0.25)
                va.format.line.fill.background()
                ch.category_axis.format.line.color.rgb = _rgb(th["muted"])
                for si, ser in enumerate(plot.series):
                    col = pal[si % len(pal)]
                    if kind == "line":
                        ser.format.line.color.rgb = _rgb(col)
                        ser.format.line.width = Pt(2.5)
                        ser.smooth = False
                        ser.marker.format.fill.solid()
                        ser.marker.format.fill.fore_color.rgb = _rgb(col)
                    elif len(plot.series) == 1:
                        for i in range(len(sl["labels"])):
                            pt = ser.points[i]
                            pt.format.fill.solid()
                            pt.format.fill.fore_color.rgb = _rgb(col if sl.get("mono") else pal[i % len(pal)])
                    else:
                        ser.format.fill.solid()
                        ser.format.fill.fore_color.rgb = _rgb(col)
                if kind in ("column", "bar"):
                    plot.gap_width = 60
            bullets_box(s, sl.get("bullets") or [], 8.8, 1.7, 3.95, 5.2, size=17)
        elif t == "table":
            heading(s, sl["title"])
            cols_n, rows_n = len(sl["columns"]), len(sl["rows"]) + 2
            tbl = s.shapes.add_table(rows_n, cols_n, Inches(0.6), Inches(1.7), Inches(12.1),
                                     Inches(min(0.55 * rows_n, 4.6))).table
            def cell(r, c, val, bold=False, fill=None, color=None, right=False):
                ce = tbl.cell(r, c)
                ce.text = str(val)
                ce.fill.solid()
                ce.fill.fore_color.rgb = _rgb(fill or th["card"])
                p = ce.text_frame.paragraphs[0]
                p.alignment = PP_ALIGN.RIGHT if right else PP_ALIGN.LEFT
                for run in p.runs:
                    run.font.size = Pt(14)
                    run.font.bold = bold
                    run.font.name = "Calibri"
                    run.font.color.rgb = _rgb(color or th["fg"])
                ce.vertical_anchor = MSO_ANCHOR.MIDDLE
            for c, h in enumerate(sl["columns"]):
                cell(0, c, h, True, th["accent"], "FFFFFF" if th["bg"] != "FFFFFF" else "FFFFFF", c > 0)
            for r, row in enumerate(sl["rows"], start=1):
                for c, v in enumerate(row):
                    cell(r, c, v, False, th["bg2"] if r % 2 == 0 else th["card"], None, c > 0)
            for c, v in enumerate(sl["total"]):
                cell(rows_n - 1, c, v, True, th["bg2"], th["accent"], c > 0)
            if sl.get("bullets"):
                text(s, sl["bullets"][0], 0.6, 1.7 + min(0.55 * rows_n, 4.6) + 0.25, 12.1, 0.6, size=14, color=th["muted"])
        elif t == "findings":
            heading(s, sl["title"])
            rect(s, 0.6, 1.7, 3.4, 5.2, th["card"], MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.05)
            sc = sl["score"]
            col = "22C55E" if sc >= 85 else "F59E0B" if sc >= 65 else "EF4444"
            text(s, str(sc), 0.6, 2.4, 3.4, 1.6, size=72, bold=True, color=col, align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE)
            text(s, "/ 100", 0.6, 3.95, 3.4, 0.5, size=18, color=th["muted"], align=PP_ALIGN.CENTER)
            text(s, str(sl["grade"]), 0.6, 4.6, 3.4, 0.6, size=22, bold=True, color=col, align=PP_ALIGN.CENTER)
            n = max(len(sl["items"]), 1)
            ih = min(1.0, 5.2 / n - 0.12)
            for i, it in enumerate(sl["items"]):
                y = 1.7 + i * (ih + 0.12)
                rect(s, 4.3, y, 8.4, ih, th["card"], MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.12)
                rect(s, 4.45, y + ih / 2 - 0.16, 0.95, 0.32, sev_color.get(it["severity"], "38BDF8"),
                     MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.5)
                text(s, it["label"], 4.45, y + ih / 2 - 0.16, 0.95, 0.32, size=11, bold=True, color="FFFFFF",
                     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                text(s, it["title"], 5.55, y + 0.05, 7.0, 0.4, size=15, bold=True)
                if it.get("detail"):
                    text(s, it["detail"], 5.55, y + 0.42, 7.0, max(ih - 0.45, 0.3), size=11, color=th["muted"])
        else:  # closing
            heading(s, sl["title"])
            items = sl.get("bullets") or []
            n = max(len(items), 1)
            ih = min(1.15, 5.2 / n - 0.15)
            for i, b in enumerate(items):
                y = 1.7 + i * (ih + 0.15)
                rect(s, 0.6, y, 12.1, ih, th["card"], MSO_SHAPE.ROUNDED_RECTANGLE, round_=0.1)
                rect(s, 0.8, y + ih / 2 - 0.3, 0.6, 0.6, pal[i % len(pal)], MSO_SHAPE.OVAL)
                text(s, str(i + 1), 0.8, y + ih / 2 - 0.3, 0.6, 0.6, size=18, bold=True, color="FFFFFF",
                     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                text(s, b, 1.65, y, 10.8, ih, size=18, anchor=MSO_ANCHOR.MIDDLE)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
