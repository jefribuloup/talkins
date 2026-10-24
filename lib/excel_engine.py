"""
Semua logic pandas/openpyxl ada di sini.
Alur: upload -> baca semua sheet -> auto-clean dasar -> simpan snapshot 'original'
      dan 'working' (identik) ke Neon. Manipulasi selanjutnya HANYA mengubah 'working'.
"""
import io
import re
import hashlib
import json
import bisect
import uuid
import math
import datetime

import pandas as pd
import numpy as np
import requests
import openpyxl
from flask import send_file

from lib import db

PAGE_SIZE = 50


# ---------- Konversi aman ke JSON ----------

def _to_jsonable(val):
    if val is None:
        return None
    if isinstance(val, float) and math.isnan(val):
        return None
    try:
        if pd.isna(val):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(val, (pd.Timestamp, datetime.datetime, datetime.date, datetime.time)):
        return val.isoformat()
    if isinstance(val, (np.datetime64,)):
        return pd.Timestamp(val).isoformat()
    if isinstance(val, (np.integer,)):
        return int(val)
    if isinstance(val, (np.floating,)):
        return float(val)
    if isinstance(val, (np.bool_,)):
        return bool(val)
    return val


def df_to_records(df: pd.DataFrame):
    records = []
    for row in df.itertuples(index=False, name=None):
        records.append([_to_jsonable(v) for v in row])
    return records


def records_to_df(records, columns):
    return pd.DataFrame(records, columns=columns)


# ---------- Grid mentah (tanpa header) ----------

def _build_grid(raw: pd.DataFrame):
    """Ubah dataframe mentah (dibaca dgn header=None) jadi grid apa adanya.

    TIDAK ada baris/kolom yang diperlakukan sebagai header: semua baris & kolom setara.
    Posisi dipertahankan persis seperti di Excel: baris 0 = baris 1 sheet, kolom 0 = kolom A,
    termasuk baris/kolom kosong di depan dan di tengah. Hanya baris/kolom kosong di
    UJUNG (setelah data terakhir) yang dipangkas. Kolom diberi nama huruf (A, B, ..., AA).
    Return None kalau sheet sama sekali kosong."""
    if raw.empty:
        return None
    mask = raw.notna().to_numpy()
    rows_any = mask.any(axis=1)
    if not rows_any.any():
        return None
    cols_any = mask.any(axis=0)
    last_r = int(np.nonzero(rows_any)[0][-1])
    last_c = int(np.nonzero(cols_any)[0][-1])
    grid = raw.iloc[: last_r + 1, : last_c + 1].copy()
    grid.columns = [_idx_to_letters(i) for i in range(grid.shape[1])]
    return grid.reset_index(drop=True)


# ---------- Deteksi kolom rumus/formula bawaan ----------

MAX_FORMULA_SCAN_ROWS = 500  # batas baris yg discan rumusnya (cukup mewakili, ga perlu scan semua)


# Referensi sel A1-style: B4, $B$4, B4:B10. Lookbehind/lookahead menghindari salah tangkap
# nama fungsi (LOG10(, ATAN2() dan referensi lintas-sheet (Sheet2!A1).
_REF_RE = re.compile(
    r"(?<![A-Za-z0-9_.!'\"])\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?(?![A-Za-z0-9_(])"
)
MAX_DEPS_PER_FORMULA = 12


def _idx_to_letters(idx):
    """0 -> A, 25 -> Z, 26 -> AA (kebalikan _letters_to_idx)."""
    n, out = idx + 1, ""
    while n > 0:
        n, rem = divmod(n - 1, 26)
        out = chr(65 + rem) + out
    return out


def _letters_to_idx(letters):
    n = 0
    for ch in letters:
        n = n * 26 + (ord(ch) - 64)
    return n - 1


def _parse_formula_deps(formula, inv_col_origin, row_origin):
    """Ekstrak 'bahan' rumus: list {"ref", "cols", "rows"} dgn kolom dipetakan ke NAMA
    kolom hasil pembersihan dan baris dipetakan ke nomor baris (1-based) di tabel
    bersih. rows=None kalau referensinya di luar area data (mis. baris judul)."""
    deps, seen = [], set()
    for m in _REF_RE.finditer(formula):
        ref = m.group(0).replace("$", "")
        if ref in seen:
            continue
        seen.add(ref)
        c1 = _letters_to_idx(m.group(1))
        r1 = int(m.group(2)) - 1
        c2 = _letters_to_idx(m.group(3)) if m.group(3) else c1
        r2 = int(m.group(4)) - 1 if m.group(4) else r1
        c_lo, c_hi = min(c1, c2), max(c1, c2)
        r_lo, r_hi = min(r1, r2), max(r1, r2)
        cols = [inv_col_origin[c] for c in range(c_lo, min(c_hi, c_lo + 60) + 1) if c in inv_col_origin]
        i_lo = bisect.bisect_left(row_origin, r_lo)
        i_hi = bisect.bisect_right(row_origin, r_hi) - 1
        rows = [i_lo + 1, i_hi + 1] if i_lo <= i_hi else None
        deps.append({"ref": ref, "cols": cols, "rows": rows})
        if len(deps) >= MAX_DEPS_PER_FORMULA:
            break
    return deps


def _detect_formulas(raw_bytes, filename, sheet_name, columns, col_origin, row_origin):
    """Baca ulang file pakai openpyxl (data_only=False) untuk mendapat 2 hal sekaligus:

    1) formula_columns: {nama_kolom: {"count","checked","ratio","sample"}} -> ringkasan
       per kolom (dipakai penanda ƒx di header, profil chat, dan fallback tab Rumus).
    2) formula_cells: {"<posisi_baris>": {nama_kolom: {"f","deps","x"}}} -> peta PER SEL, hanya
       sel yang benar-benar rumus (sparse). posisi_baris = index 0-based baris di
       working_data SAAT INGEST. Kalau nanti baris di-sort/filter/dedupe, posisi itu
       bergeser -> ditandai lewat row_order_dirty di db, bukan dihapus.

    Kenapa baca ulang: pandas hanya membaca HASIL hitung rumus (angka), bukan
    string rumusnya. Hanya .xlsx/.xlsm; format .xls lama (xlrd) tidak bisa.
    Mengembalikan (formula_columns, formula_cells).
    """
    if not filename.lower().endswith((".xlsx", ".xlsm")):
        return {}, {}
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=False)
    except Exception:
        return {}, {}
    if sheet_name not in wb.sheetnames:
        wb.close()
        return {}, {}
    ws = wb[sheet_name]

    scan_rows = list(enumerate(row_origin[:MAX_FORMULA_SCAN_ROWS]))  # (posisi, baris_asli)
    inv_col_origin = {idx: name for name, idx in col_origin.items()}
    formula_columns = {}
    formula_cells = {}

    for col_name in columns:
        orig_col = col_origin.get(col_name)
        if orig_col is None:
            continue
        checked = formula_hits = 0
        sample_formula = None
        for row_pos, orig_row in scan_rows:
            # +1 krn openpyxl 1-based, sedangkan col_origin/row_origin 0-based
            val = ws.cell(row=orig_row + 1, column=orig_col + 1).value
            if val is None:
                continue
            checked += 1
            if isinstance(val, str) and val.startswith("="):
                formula_hits += 1
                if sample_formula is None:
                    sample_formula = val
                formula_cells.setdefault(str(row_pos), {})[col_name] = {
                    "f": val,
                    "deps": _parse_formula_deps(val, inv_col_origin, row_origin),
                    "x": "!" in val,  # merujuk sheet lain
                }
        if formula_hits:
            formula_columns[col_name] = {
                "count": formula_hits,
                "checked": checked,
                "ratio": round(formula_hits / checked, 2) if checked else 0,
                "sample": sample_formula,
            }

    wb.close()
    return formula_columns, formula_cells


# ---------- Ingest ----------

def _pick_engine(filename: str) -> str:
    lower = filename.lower()
    if lower.endswith(".xls"):
        return "xlrd"
    return "openpyxl"


def ingest_file(file_storage, session_id):
    filename = file_storage.filename
    if not filename.lower().endswith((".xlsx", ".xls", ".xlsm")):
        raise ValueError("Format harus .xlsx / .xls / .xlsm")

    raw_bytes = file_storage.read()
    return _ingest_bytes(raw_bytes, filename, session_id, source="upload")


# ---------- Ingest dari link spreadsheet (mis. Google Sheets publik) ----------

_GOOGLE_SHEET_ID_RE = re.compile(r"/spreadsheets/d/([a-zA-Z0-9-_]+)")


def _resolve_spreadsheet_url(url: str):
    """Kalau link Google Sheets biasa -> ubah jadi link export .xlsx langsung.
    Link lain (mis. URL download .xlsx dari cloud storage) dipakai apa adanya."""
    m = _GOOGLE_SHEET_ID_RE.search(url)
    if m:
        sheet_id = m.group(1)
        return (
            f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx",
            f"google-sheet-{sheet_id}.xlsx",
        )
    name = url.split("?")[0].rstrip("/").split("/")[-1] or "spreadsheet.xlsx"
    if not name.lower().endswith((".xlsx", ".xls", ".xlsm")):
        name += ".xlsx"
    return url, name


def ingest_from_url(url, session_id):
    if not url or not url.strip().lower().startswith(("http://", "https://")):
        raise ValueError("Link tidak valid")

    fetch_url, filename = _resolve_spreadsheet_url(url.strip())

    # Link yang sama sudah pernah diproses di sesi ini -> jangan unduh & buat salinan lagi
    dup = db.find_duplicate(session_id, source_key=fetch_url)
    if dup:
        return _duplicate_response(dup)

    try:
        resp = requests.get(fetch_url, timeout=20, allow_redirects=True)
        resp.raise_for_status()
    except requests.RequestException:
        raise ValueError(
            "Gagal mengambil data dari link. Pastikan link bisa diakses publik "
            "(Google Sheets: Share -> Anyone with the link -> Viewer)."
        )

    return _ingest_bytes(resp.content, filename, session_id, source="url", source_key=fetch_url)


# ---------- Logic inti ingest, dipakai baik oleh upload file maupun link ----------

def _duplicate_response(dup):
    """Bentuk respons standar kalau file/link yang sama sudah pernah diunggah.
    Frontend yang menanyakan ke user apakah mau membuka file yang sudah ada."""
    return {
        "duplicate": True,
        "file_id": dup["id"],
        "filename": dup["filename"],
        "sheets": dup.get("sheet_names") or [],
        "uploaded_at": dup["uploaded_at"],
        "source": dup.get("source", "upload"),
    }


def _ingest_bytes(raw_bytes, filename, session_id, source="upload", source_key=None):
    # Cegah dobel: isi file identik (SHA-256) dgn yang sudah ada di sesi ini
    content_hash = hashlib.sha256(raw_bytes).hexdigest()
    dup = db.find_duplicate(session_id, content_hash=content_hash, source_key=source_key)
    if dup:
        return _duplicate_response(dup)

    engine = _pick_engine(filename)
    try:
        xls = pd.ExcelFile(io.BytesIO(raw_bytes), engine=engine)
    except Exception:
        raise ValueError(
            "File/spreadsheet tidak bisa dibaca. Pastikan formatnya didukung "
            "(.xlsx/.xls/.xlsm) dan, untuk link, aksesnya publik."
        )
    file_id = str(uuid.uuid4())
    db.create_file(file_id, session_id, filename, [], source=source,
                   content_hash=content_hash, source_key=source_key)

    sheet_names = []
    for sheet_name in xls.sheet_names:
        # dtype=object: nilai dibaca apa adanya (angka tetap angka, teks tetap teks)
        raw_df = pd.read_excel(xls, sheet_name=sheet_name, header=None, dtype=object)
        grid = _build_grid(raw_df)
        if grid is None:
            continue
        columns = list(grid.columns)  # ["A", "B", ...] = huruf kolom Excel asli
        records = df_to_records(grid)
        # Grid = sheet asli (tanpa pemangkasan di depan) -> posisi kolom/baris identik dgn Excel
        col_origin = {name: i for i, name in enumerate(columns)}
        row_origin = list(range(len(records)))
        formula_columns, formula_cells = _detect_formulas(
            raw_bytes, filename, sheet_name, columns, col_origin, row_origin
        )
        db.save_sheet(file_id, sheet_name, records, columns, formula_columns, formula_cells)
        sheet_names.append(sheet_name)

    if not sheet_names:
        db.delete_file(session_id, file_id)
        raise ValueError("Tidak ada data yang bisa dibaca dari file/link ini")

    db.update_file_sheets(file_id, sheet_names)
    return {
        "file_id": file_id,
        "filename": filename,
        "sheets": sheet_names,
        "source": source,
    }


# ---------- Preview & metadata ----------

def get_sheet_list(file_id):
    return {"sheets": db.list_sheet_names(file_id)}


# ---------- Pemulihan file lama (diunggah sebelum grid tanpa header) ----------
# Di file lama, baris header disimpan sebagai NAMA kolom (bukan data), jadi tidak ikut tampil.
# Saat sheet dibuka, sheet lama otomatis dikonversi: header dikembalikan jadi baris biasa di
# posisi aslinya. Penanda file lama = original_meta punya kunci "header_row".

_LEGACY_AUTO_NAME = re.compile(r"^kolom_\d+$")
_NUMERIC_TEXT = re.compile(r"^-?(0|[1-9]\d*)(\.\d+)?$")


def _legacy_header_value(name):
    if not isinstance(name, str) or _LEGACY_AUTO_NAME.match(name):
        return None
    n = re.sub(r"^(-?\d+)\.0+$", r"\1", name)
    if _NUMERIC_TEXT.match(n):
        return float(n) if "." in n else int(n)
    return name


def _legacy_to_grid(records, columns, letters, header_row):
    """records/columns versi lama -> (grid baris-baris, daftar huruf kolom, peta nama->huruf)."""
    idx = {}
    for i, c in enumerate(columns):
        idx[c] = _letters_to_idx(letters[c]) if c in letters else i
    width = (max(idx.values()) + 1) if idx else 0
    header_pos = max(0, int(header_row or 1) - 1)

    def blank():
        return [None] * width

    grid = [blank() for _ in range(header_pos)]
    head = blank()
    for c in columns:
        head[idx[c]] = _legacy_header_value(c)
    grid.append(head)
    for rec in records:
        row = blank()
        for c, v in zip(columns, rec):
            row[idx[c]] = v
        grid.append(row)
    return grid, _letters_list(width), {c: _idx_to_letters(i) for c, i in idx.items()}


def _migrate_legacy_sheet(file_id, sheet_name):
    conn = db.get_conn()
    cur = conn.cursor()
    cur.execute(
        "SELECT original_data, working_data, columns, original_meta, formula_columns, column_letters "
        "FROM sheets WHERE file_id=%s AND sheet_name=%s",
        (file_id, sheet_name),
    )
    row = cur.fetchone()
    if not row or "header_row" not in (row[3] or {}):
        cur.close()
        conn.close()
        return
    orig_data, work_data, work_cols, meta, work_fcols, work_letters = row
    header_row = meta.get("header_row", 1)
    o_grid, o_cols, o_map = _legacy_to_grid(
        orig_data, meta.get("columns") or work_cols, meta.get("column_letters") or {}, header_row)
    w_grid, w_cols, w_map = _legacy_to_grid(
        work_data, work_cols, work_letters or {}, header_row)

    def remap(fc, mapping):
        return {mapping[k]: v for k, v in (fc or {}).items() if k in mapping}

    new_meta = {
        "columns": o_cols,
        "formula_columns": remap(meta.get("formula_columns"), o_map),
        "formula_cells": {},  # peta rumus per-sel versi lama tidak bisa dipetakan -> upload ulang untuk detailnya
        "column_letters": {},
    }
    cur.execute(
        "UPDATE sheets SET original_data=%s, working_data=%s, columns=%s, formula_columns=%s, "
        "formula_cells='{}', row_order_dirty=false, column_letters='{}', original_meta=%s "
        "WHERE file_id=%s AND sheet_name=%s",
        (json.dumps(o_grid), json.dumps(w_grid), json.dumps(w_cols),
         json.dumps(remap(work_fcols, w_map)), json.dumps(new_meta), file_id, sheet_name),
    )
    conn.commit()
    cur.close()
    conn.close()


def get_recap(file_id, sheet_name, lang="id"):
    """Rekapitulasi (duplikat, format, error, analisa) untuk data sheet saat ini."""
    from lib import recap_engine
    _migrate_legacy_sheet(file_id, sheet_name)
    records = db.get_sheet(file_id, sheet_name, version="working")
    if records is None:
        raise ValueError("Sheet tidak ditemukan")
    columns = _get_columns(file_id, sheet_name)
    formula_columns, formula_cells, dirty, _ = db.get_formula_info(file_id, sheet_name)
    return recap_engine.build_recap(
        records, columns, formula_columns, {} if dirty else formula_cells, lang)


MAX_ROWS_PER_REQUEST = 500
MAX_COL_CHARS = 60


def _col_chars(records, ncols):
    """Panjang teks terpanjang tiap kolom (jumlah karakter tampilan, maks MAX_COL_CHARS).
    0 = kolom tanpa isi -> frontend memakai lebar standar. Dipakai untuk lebar kolom dinamis."""
    widths = [0] * ncols
    for row in records:
        for j, v in enumerate(row[:ncols]):
            if v is None:
                continue
            if isinstance(v, str):
                n = len(v)
            elif isinstance(v, float):
                if v != v:
                    continue
                n = len(str(int(v))) if v == int(v) else len(str(v))
            else:
                n = len(str(v))
            if n > widths[j]:
                widths[j] = min(n, MAX_COL_CHARS)
    return widths



def preview_sheet(file_id, sheet_name, row_start=0, row_count=200, with_meta=True):
    """Kirim potongan baris [row_start, row_start+row_count) dari grid (tanpa header).

    Frontend memuat grid per blok saat user scroll. with_meta=True menambahkan info
    ringkas sheet (jumlah baris/kolom, info rumus) -> cukup diminta sekali per sheet."""
    if with_meta:
        _migrate_legacy_sheet(file_id, sheet_name)
    records = db.get_sheet(file_id, sheet_name, version="working")
    if records is None:
        raise ValueError("Sheet tidak ditemukan")
    total_rows = len(records)
    row_start = max(0, min(int(row_start), total_rows))
    row_count = max(1, min(int(row_count), MAX_ROWS_PER_REQUEST))
    out = {
        "row_start": row_start,
        "rows": records[row_start:row_start + row_count],
    }
    if with_meta:
        columns = _get_columns(file_id, sheet_name)
        formula_columns, formula_cells, row_order_dirty, _ = db.get_formula_info(file_id, sheet_name)
        out.update({
            "columns": columns,
            "total_rows": total_rows,
            "total_cols": len(columns),
            "col_chars": _col_chars(records, len(columns)),
            "formula_columns": formula_columns,
            # peta rumus per sel (sparse, hanya baris yg discan saat ingest) -> key = indeks baris
            "formula_cells": {} if row_order_dirty else formula_cells,
            "row_order_dirty": row_order_dirty,
        })
    return out


def _get_columns(file_id, sheet_name):
    conn = db.get_conn()
    cur = conn.cursor()
    cur.execute(
        "SELECT columns FROM sheets WHERE file_id=%s AND sheet_name=%s",
        (file_id, sheet_name),
    )
    row = cur.fetchone()
    cur.close()
    conn.close()
    return row[0] if row else []


# ---------- Manipulasi working copy ----------

def _letters_list(n):
    return [_idx_to_letters(i) for i in range(n)]


def _norm_col(c):
    return str(c).strip().upper()


def apply_operation(file_id, op):
    sheet_name = op.get("sheet")
    action = op.get("action")
    columns = _get_columns(file_id, sheet_name)
    records = db.get_sheet(file_id, sheet_name, version="working")
    df = records_to_df(records, columns)

    # Kolom dirujuk dengan huruf Excel (A, B, ...), tidak ada "nama kolom"/header.
    if op.get("column") is not None:
        op["column"] = _norm_col(op["column"])
    if op.get("columns") is not None:
        op["columns"] = [_norm_col(c) for c in op["columns"] if str(c).strip()]
    for c in ([op["column"]] if op.get("column") else []) + (op.get("columns") or []):
        if c not in df.columns:
            raise ValueError(f"Kolom {c} tidak ada di sheet ini")

    rename_map = None  # {huruf_lama: huruf_baru}, hanya diisi saat kolom dihapus

    if action == "drop_na":
        cols = op.get("columns") or None
        df = df.dropna(subset=cols) if cols else df.dropna(how="all")

    elif action == "fill_na":
        value = op.get("value", "")
        cols = op.get("columns") or df.columns.tolist()
        df[cols] = df[cols].fillna(value)

    elif action == "drop_duplicates":
        cols = op.get("columns") or None
        df = df.drop_duplicates(subset=cols)

    elif action == "drop_column":
        drop = set(op.get("columns", []))
        keep = [c for c in df.columns if c not in drop]
        df = df[keep]
        rename_map = {old: _idx_to_letters(i) for i, old in enumerate(keep)}

    elif action == "filter_rows":
        col, cond, val = op["column"], op["condition"], op["value"]
        series = df[col]
        if cond == "eq":
            df = df[series == val]
        elif cond == "neq":
            df = df[series != val]
        elif cond == "gt":
            df = df[pd.to_numeric(series, errors="coerce") > float(val)]
        elif cond == "lt":
            df = df[pd.to_numeric(series, errors="coerce") < float(val)]
        elif cond == "contains":
            df = df[series.astype(str).str.contains(str(val), case=False, na=False)]

    elif action == "sort":
        col = op["column"]
        ascending = op.get("ascending", True)
        df = df.sort_values(by=col, ascending=ascending)

    elif action == "cast_numeric":
        cols = op.get("columns", [])
        for c in cols:
            df[c] = pd.to_numeric(df[c], errors="coerce")

    else:
        raise ValueError(f"Aksi tidak dikenal: {action}")

    df = df.reset_index(drop=True)
    # Nama kolom selalu = huruf berurutan dari A (seperti Excel setelah kolom dihapus)
    df.columns = _letters_list(len(df.columns))
    new_columns = list(df.columns)
    new_records = df_to_records(df)
    db.update_working_data(file_id, sheet_name, new_records, new_columns)

    # Aksi yang mengubah urutan/jumlah baris atau menggeser kolom -> peta rumus per-sel
    # (berbasis posisi saat ingest) tidak akurat lagi sampai user reset ke data asli.
    if action in ("drop_na", "drop_duplicates", "filter_rows", "sort", "drop_column"):
        db.set_row_order_dirty(file_id, sheet_name, True)

    # Kolom dihapus -> huruf kolom di info rumus ikut bergeser.
    if rename_map is not None:
        formula_columns = db.get_formula_columns(file_id, sheet_name)
        db.update_formula_columns(
            file_id, sheet_name,
            {rename_map[k]: v for k, v in formula_columns.items() if k in rename_map},
        )

    return {
        "columns": new_columns,
        "row_count": len(new_records),
        "preview": new_records[:PAGE_SIZE],
    }


def reset_to_original(file_id):
    db.reset_working_data(file_id)


# ---------- Export working copy jadi file baru (bukan menimpa file asli) ----------

def export_working_copy(file_id):
    file_meta = db.get_file(file_id)
    sheet_names = db.list_sheet_names(file_id)
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for sheet_name in sheet_names:
            columns = _get_columns(file_id, sheet_name)
            records = db.get_sheet(file_id, sheet_name, version="working")
            df = records_to_df(records, columns)
            # header=False: grid ditulis apa adanya (baris 1 = baris 1), tanpa baris huruf A/B/C
            df.to_excel(writer, sheet_name=sheet_name[:31], index=False, header=False)
    output.seek(0)
    download_name = f"edited_{file_meta['filename']}"
    return send_file(
        output,
        as_attachment=True,
        download_name=download_name,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
